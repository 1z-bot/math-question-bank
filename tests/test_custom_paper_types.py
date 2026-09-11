"""Regression for selected custom types disappearing from ordinary papers."""

import io
import os
from pathlib import Path
import re
import shutil
import subprocess
import zipfile

from docx import Document
import pytest

from mathbank.paper_helper import build_latex_document, compile_tex_to_pdf
from mathbank.question_types import normalize_section_order, paper_type_order
from mathbank.word_export_helper import build_word_document


TYPES = [
    {"value": "proof", "label": "证明题"},
    {"value": "calculation", "label": "计算题"},
    {"value": "proof", "label": "重复配置不生成重复大题"},
]


def question(qid, kind, score=5):
    return {
        "question": {
            "id": qid, "question_type": kind,
            "content": f"STEM{qid} 求下式的值。",
            "answer_markdown": f"ANSWER{qid} 演算结果。",
        },
        "score": score, "solution_space": "1.5",
    }


@pytest.mark.parametrize("paper_type", ["quiz", "exam"])
@pytest.mark.parametrize("mixed", [False, True])
@pytest.mark.parametrize("answers", [False, True])
def test_custom_questions_labels_order_scores_and_answers(paper_type, mixed, answers):
    data = [question(196, "calculation"), question(188, "calculation", 7)]
    expected = [196, 188]
    if mixed:
        data += [question(3, "proof"), question(4, "single_choice"), question(5, "历史题型")]
        expected = [4, 3, 196, 188, 5]
    tex = build_latex_document("题型回归", "", paper_type, data, include_answers=answers, question_types=TYPES)
    docx, diagnostics = build_word_document("题型回归", "", paper_type, data, include_answers=answers, question_types=TYPES)
    document = Document(io.BytesIO(docx))
    text = "\n".join(p.text for p in document.paragraphs)
    for output in [tex, text]:
        assert [int(n) for n in re.findall(r"STEM(\d+)", output)] == expected
        assert [int(n) for n in re.findall(r"ANSWER(\d+)", output)] == (expected if answers else [])
        assert "计算题" in output
        if mixed:
            assert "证明题" in output and "历史题型" in output
        if paper_type == "exam":
            assert "共 12 分" in output
    assert r"\begin{problem}[points = 7]" in tex
    assert (r"\vspace*{1.5cm}" in tex) is (not answers)
    assert [int(n) for n in re.findall(r"^(\d+)\. STEM", text, re.M)] == list(range(1, len(data) + 1))
    assert diagnostics["failed_formulas"] == 0
    assert any(p.paragraph_format.space_after and p.paragraph_format.space_after.cm > 1 for p in document.paragraphs)


def test_exam_19_keeps_existing_builtin_sections_and_fixed_numbers():
    data = [question(i, kind) for i, kind in enumerate(["single_choice", "multi_choice", "fill_in_blank", "detailed_answer"], 1)]
    reversed_order = ["detailed_answer", "fill_in_blank", "multi_choice", "single_choice", "calculation"]
    tex = build_latex_document("高考", "", "exam_19", data, question_types=TYPES, section_order=reversed_order)
    assert tex == build_latex_document("高考", "", "exam_19", data)
    assert paper_type_order(["proof", "calculation"], "exam_19", TYPES) == ["single_choice", "multi_choice", "fill_in_blank", "detailed_answer"]
    docx, _ = build_word_document("高考", "", "exam_19", data, question_types=TYPES, section_order=reversed_order)
    text = "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)
    assert re.findall(r"^(\d+)\. STEM", text, re.M) == ["1", "9", "12", "15"]


def test_custom_label_is_literal_text_in_tex_and_word():
    label = r"计算_提高 & 证明 50% #1 {A} $B$ \C ^ ~"
    types = [{"value": "calculation", "label": label}]
    data = [question(196, "calculation")]
    tex = build_latex_document("题型回归", "", "quiz", data, question_types=types)
    assert r"计算\_提高 \& 证明 50\% \#1 \{A\} \textdollar{}B\textdollar{} \textbackslash{}C" in tex
    docx, _ = build_word_document("题型回归", "", "quiz", data, question_types=types)
    assert label in [p.text for p in Document(io.BytesIO(docx)).paragraphs]


def test_frontend_type_order_matches_backend_including_old_numeric_and_prototype_keys():
    source = (Path(__file__).resolve().parents[1] / "static/js/paper.js").read_text()
    start = source.index("    function isWrittenQuestionType(")
    end = source.index("    function getDifficultyBadge(", start)
    script = "const window = {systemMetadata: {question_types: [{value:'proof',label:'证明题'}]}};\n" + source[start:end]
    script += """
const present = ['calculation','20','3','__proto__','proof','calculation'];
const actual = getPaperTypeOrder(present, 'quiz');
const expected = ['single_choice','multi_choice','fill_in_blank','detailed_answer','proof','calculation','20','3','__proto__'];
if (JSON.stringify(actual) !== JSON.stringify(expected)) throw new Error(actual);
if (getQuestionTypeCn('__proto__') !== '__proto__') throw new Error('prototype label');
if (!isWrittenQuestionType('calculation') || isWrittenQuestionType('fill_in_blank')) throw new Error('layout');
if (getPaperTypeOrder(present, 'exam_19').length !== 4) throw new Error('exam19');
"""
    result = subprocess.run([shutil.which("node"), "-e", script], capture_output=True, text=True, timeout=10)
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("paper_type", ["quiz", "exam"])
@pytest.mark.parametrize("endpoint", ["tex", "word", "pdf", "bundle"])
def test_export_endpoints_pass_server_custom_labels(client, db_session, monkeypatch, paper_type, endpoint):
    import main
    from mathbank.database import Question

    q = Question(content="ENDPOINTSTEM", question_type="calculation", difficulty="easy")
    choice = Question(content="CHOICESTEM", question_type="single_choice", difficulty="easy")
    db_session.add_all([q, choice])
    db_session.commit()
    monkeypatch.setitem(main.METADATA_CACHE, "question_types", TYPES)
    seen = []

    def fake_compile(tex, *args, **kwargs):
        seen.append(tex)
        assert "计算题" in tex and "ENDPOINTSTEM" in tex
        assert tex.index("ENDPOINTSTEM") < tex.index("CHOICESTEM")
        return b"%PDF-1.7\nfixture", ""

    monkeypatch.setattr(main, "compile_tex_to_pdf", fake_compile)
    response = client.post(f"/api/paper/export/{endpoint}", headers={"X-Local-Token": main.LOCAL_TOKEN}, json={
        "title": "自定义题型", "paper_type": paper_type,
        "section_order": ["calculation", "single_choice"],
        "questions": [{"id": choice.id, "score": 5}, {"id": q.id, "score": 5, "solution_space": "0"}],
    })
    assert response.status_code == 200, response.text
    if endpoint in ["tex", "word", "bundle"]:
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            outputs = []
            for name in archive.namelist():
                if name.endswith(".tex"):
                    outputs.append(archive.read(name).decode())
                elif name.endswith(".docx"):
                    outputs.append("\n".join(p.text for p in Document(io.BytesIO(archive.read(name))).paragraphs))
            assert len(outputs) == 2
            assert all("计算题" in output and "ENDPOINTSTEM" in output for output in outputs)
            assert all(output.index("ENDPOINTSTEM") < output.index("CHOICESTEM") for output in outputs)
    if endpoint in ["pdf", "bundle"]:
        assert seen


@pytest.mark.parametrize("paper_type", ["exam", "quiz"])
def test_manual_section_order_controls_body_and_answer_numbering(paper_type):
    data = [question(1, "single_choice"), question(2, "calculation"), question(3, "proof"), question(4, "calculation")]
    order = ["proof", "calculation", "proof", "single_choice"]
    tex = build_latex_document("顺序", "", paper_type, data, include_answers=True, question_types=TYPES, section_order=order)
    docx, _ = build_word_document("顺序", "", paper_type, data, include_answers=True, question_types=TYPES, section_order=order)
    text = "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)
    for output in [tex, text]:
        assert re.findall(r"STEM(\d+)", output) == ["3", "2", "4", "1"]
        assert re.findall(r"ANSWER(\d+)", output) == ["3", "2", "4", "1"]
    assert re.findall(r"^(\d+)\. STEM", text, re.M) == ["1", "2", "3", "4"]


def test_saved_section_order_roundtrip_and_legacy_default(client, db_session):
    import main
    from mathbank.database import Paper, Question

    q = Question(content="顺序保存测试", question_type="calculation", difficulty="easy")
    db_session.add(q)
    db_session.commit()
    response = client.post("/api/paper/save", headers={"X-Local-Token": main.LOCAL_TOKEN}, json={
        "title": "顺序保存", "paper_type": "quiz", "section_order": ["calculation", "single_choice", "proof"],
        "questions": [{"id": q.id, "score": 5}],
    })
    assert response.status_code == 200
    detail = client.get(f"/api/papers/{response.json()['paper_id']}").json()["data"]
    assert detail["section_order"] == ["calculation", "single_choice", "proof"]
    assert detail["questions"][0]["question"]["question_type"] == "calculation"
    assert Paper(title="旧试卷", metadata_json="{}").to_dict()["section_order"] == []
    assert Paper(title="旧试卷", metadata_json='{"section_order": "bad"}').to_dict()["section_order"] == []


def test_partial_and_malformed_order_cannot_drop_new_types():
    present = ["proof", "calculation", "历史题型"]
    order = ["calculation", None, ["proof"], "calculation", "", "single_choice"]
    assert normalize_section_order(order) == ["calculation", "single_choice"]
    assert paper_type_order(present, "quiz", TYPES, order) == [
        "calculation", "single_choice", "multi_choice", "fill_in_blank", "detailed_answer", "proof", "历史题型",
    ]
    assert paper_type_order(present, "quiz", TYPES, "broken") == paper_type_order(present, "quiz", TYPES)


def test_frontend_section_move_preserves_questions_hidden_positions_and_snapshot():
    source = (Path(__file__).resolve().parents[1] / "static/js/paper.js").read_text()
    helpers = source[source.index("    function isWrittenQuestionType("):source.index("    function getDifficultyBadge(")]
    handler = source[source.index("    window.movePaperSection ="):source.index("    // Reorder Items strictly")]
    signature = source[source.index("    function getPaperCartSignature("):source.index("    function beginPaperAction(")]
    script = """
const assert = require('assert');
const window = {systemMetadata:{question_types:[]}, PaperStore:{
    cart:[{id:1},{id:2},{id:3}], meta:{paper_type:'quiz',section_order:[]},
    questionsMap:{1:{content:'A',question_type:'single_choice'},2:{content:'B',question_type:'calculation'},3:{content:'C',question_type:'proof'}}
}};
const document = {querySelectorAll:()=>[]};
let saved;
function saveMetaToStorage(){saved=JSON.parse(JSON.stringify(window.PaperStore.meta));}
window.renderPaperCanvas = ()=>{};
""" + helpers + handler + signature + """
const originalCart=JSON.stringify(window.PaperStore.cart);
const before=getPaperCartSignature();
window.movePaperSection('calculation','up');
const order=window.PaperStore.meta.section_order;
assert(order.indexOf('calculation')<order.indexOf('single_choice'));
assert.deepEqual(order,['calculation','single_choice','multi_choice','fill_in_blank','detailed_answer','proof']);
assert.notEqual(before,getPaperCartSignature());
assert.equal(JSON.stringify(window.PaperStore.cart),originalCart);
assert.deepEqual(saved.section_order,order);
// Remove/re-add a section without rewriting its saved position.
assert(getPaperTypeOrder(['single_choice','proof'],'quiz',order).indexOf('calculation')<getPaperTypeOrder(['single_choice','proof'],'quiz',order).indexOf('single_choice'));
const current=getPaperCartSignature();
window.movePaperSection('calculation','up');
assert.equal(current,getPaperCartSignature());
window.PaperStore.meta.paper_type='exam_19';
window.movePaperSection('proof','up');
assert.deepEqual(window.PaperStore.meta.section_order,order);
assert.deepEqual(getPaperTypeOrder(['proof'],'exam_19',order),['single_choice','multi_choice','fill_in_blank','detailed_answer']);
window.PaperStore.meta.paper_type='exam';
window.movePaperSection('calculation','down');
assert(window.PaperStore.meta.section_order.indexOf('single_choice')<window.PaperStore.meta.section_order.indexOf('calculation'));
"""
    result = subprocess.run([shutil.which("node"), "-e", script], capture_output=True, text=True, timeout=10)
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(os.environ.get("MATHBANK_TEST_CUSTOM_TYPES_NATIVE") != "1" or not shutil.which("xelatex"), reason="Opt-in native PDF regression")
@pytest.mark.parametrize("paper_type", ["quiz", "exam"])
def test_native_pdf_custom_questions_are_visible(paper_type):
    import pymupdf as fitz

    data = [question(196, "calculation"), question(188, "proof")]
    tex = build_latex_document("自定义题型回归", "", paper_type, data, include_answers=True, question_types=TYPES, section_order=["calculation", "proof"])
    pdf, diagnostic = compile_tex_to_pdf(tex)
    assert pdf, diagnostic
    with fitz.open(stream=pdf, filetype="pdf") as document:
        text = "".join(page.get_text() for page in document)
    for expected in ["计算题", "证明题", "STEM196", "STEM188", "ANSWER196", "ANSWER188"]:
        assert expected in text
    assert text.index("STEM196") < text.index("STEM188")
